from docx import Document
import datetime
import os

TEMPLATE_PATH = 'templates/decision_report_template.docx'
REPORT_DIR = 'reports'

def generate_report(customer_id, customer_data, analysis_data):
    """
    生成决策报告
    
    参数：
    customer_data 索引：
    0:id, 1:name, 2:phone, 3:birth_date, 4:birth_hour, 5:gender,
    6:decision_type, 7:decision_desc, 8:background, 9:status, 10:created_at
    
    analysis_data 索引：
    0:id, 1:customer_id, 2:system_state, 3:core_contradiction, 4:opportunities, 5:risks,
    6:path_1_name, 7:path_1_prob, 8:path_1_risk, 9:path_1_analysis,
    10:path_2_name, 11:path_2_prob, 12:path_2_risk, 13:path_2_analysis,
    14:path_3_name, 15:path_3_prob, 16:path_3_risk, 17:path_3_analysis,
    18:final_advice_short, 19:final_advice_medium, 20:final_advice_long,
    21:key_warning, 22:analyst_name, 23:updated_at
    """
    os.makedirs(REPORT_DIR, exist_ok=True)
    
    if not os.path.exists(TEMPLATE_PATH):
        create_default_template()
    
    doc = Document(TEMPLATE_PATH)
    
    data = {
        # 客户基本信息
        'customer_name': customer_data[1] or '客户',
        'customer_phone': customer_data[2] or '未提供',
        'customer_birth_date': customer_data[3] or '未提供',
        'customer_birth_hour': customer_data[4] or '未提供',
        'customer_gender': customer_data[5] or '未提供',
        'customer_decision_type': customer_data[6] or '未填写',
        'customer_decision_desc': customer_data[7] or '未填写',
        'customer_background': customer_data[8] or '未提供',
        'report_date': datetime.datetime.now().strftime('%Y年%m月%d日'),
        
        # 分析结论
        'analyst_name': analysis_data[22] if analysis_data else '分析师',
        'system_state': analysis_data[2] if analysis_data else '待评估',
        'core_contradiction': analysis_data[3] if analysis_data else '待分析',
        'opportunities': analysis_data[4] if analysis_data else '待分析',
        'risks': analysis_data[5] if analysis_data else '待分析',
        
        'path_1_name': analysis_data[6] if analysis_data else '',
        'path_1_prob': analysis_data[7] if analysis_data else '',
        'path_1_risk': analysis_data[8] if analysis_data else '',
        'path_1_analysis': analysis_data[9] if analysis_data else '',
        
        'path_2_name': analysis_data[10] if analysis_data else '',
        'path_2_prob': analysis_data[11] if analysis_data else '',
        'path_2_risk': analysis_data[12] if analysis_data else '',
        'path_2_analysis': analysis_data[13] if analysis_data else '',
        
        'path_3_name': analysis_data[14] if analysis_data else '',
        'path_3_prob': analysis_data[15] if analysis_data else '',
        'path_3_risk': analysis_data[16] if analysis_data else '',
        'path_3_analysis': analysis_data[17] if analysis_data else '',
        
        'final_advice_short': analysis_data[18] if analysis_data else '',
        'final_advice_medium': analysis_data[19] if analysis_data else '',
        'final_advice_long': analysis_data[20] if analysis_data else '',
        'key_warning': analysis_data[21] if analysis_data else '无'
    }
    
    # 替换段落中的占位符
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = f'{{{{{key}}}}}'
            if placeholder in paragraph.text:
                for run in paragraph.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, str(value))
    
    # 替换表格中的占位符
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in data.items():
                        placeholder = f'{{{{{key}}}}}'
                        if placeholder in paragraph.text:
                            for run in paragraph.runs:
                                if placeholder in run.text:
                                    run.text = run.text.replace(placeholder, str(value))
    
    filename = f"{REPORT_DIR}/{customer_data[1] or '客户'}_决策报告_{datetime.datetime.now().strftime('%Y%m%d')}.docx"
    doc.save(filename)
    return filename


def create_default_template():
    """创建默认的Word报告模板（含客户基本信息区域）"""
    os.makedirs('templates', exist_ok=True)
    doc = Document()
    
    doc.add_heading('人生决策辅助报告', 0)
    doc.add_paragraph('')
    
    # 客户信息确认
    doc.add_heading('📋 客户信息确认', level=1)
    doc.add_paragraph('请核对以下信息是否为您本人：')
    doc.add_paragraph('')
    
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = '姓名'
    table.cell(0, 1).text = '{{customer_name}}'
    table.cell(1, 0).text = '手机号'
    table.cell(1, 1).text = '{{customer_phone}}'
    table.cell(2, 0).text = '性别'
    table.cell(2, 1).text = '{{customer_gender}}'
    table.cell(3, 0).text = '出生日期'
    table.cell(3, 1).text = '{{customer_birth_date}}  {{customer_birth_hour}}'
    table.cell(4, 0).text = '决策问题'
    table.cell(4, 1).text = '{{customer_decision_desc}}'
    
    doc.add_paragraph('')
    doc.add_paragraph(f'报告日期：{{{{report_date}}}}')
    doc.add_paragraph(f'分析师：{{{{analyst_name}}}}')
    doc.add_paragraph('')
    
    # 正文
    doc.add_heading('一、系统状态扫描', level=1)
    doc.add_paragraph(f'当前状态：{{{{system_state}}}}')
    doc.add_paragraph(f'核心矛盾：{{{{core_contradiction}}}}')
    doc.add_paragraph(f'机会点：{{{{opportunities}}}}')
    doc.add_paragraph(f'风险点：{{{{risks}}}}')
    doc.add_paragraph('')
    
    doc.add_heading('二、决策路径推演', level=1)
    doc.add_paragraph(f'路径一：{{{{path_1_name}}}}（概率{{{{path_1_prob}}}}%，风险{{{{path_1_risk}}}}）')
    doc.add_paragraph(f'分析：{{{{path_1_analysis}}}}')
    doc.add_paragraph('')
    doc.add_paragraph(f'路径二：{{{{path_2_name}}}}（概率{{{{path_2_prob}}}}%，风险{{{{path_2_risk}}}}）')
    doc.add_paragraph(f'分析：{{{{path_2_analysis}}}}')
    doc.add_paragraph('')
    doc.add_paragraph(f'路径三：{{{{path_3_name}}}}（概率{{{{path_3_prob}}}}%，风险{{{{path_3_risk}}}}）')
    doc.add_paragraph(f'分析：{{{{path_3_analysis}}}}')
    doc.add_paragraph('')
    
    doc.add_heading('三、执行建议', level=1)
    doc.add_paragraph(f'短期建议（1个月）：{{{{final_advice_short}}}}')
    doc.add_paragraph(f'中期建议（3个月）：{{{{final_advice_medium}}}}')
    doc.add_paragraph(f'长期建议（1年）：{{{{final_advice_long}}}}')
    doc.add_paragraph('')
    
    doc.add_heading('四、风险提示', level=1)
    doc.add_paragraph(f'⚠️ {{{{key_warning}}}}')
    doc.add_paragraph('')
    doc.add_paragraph('---')
    doc.add_paragraph('免责声明：本报告基于传统文化分析模型生成，仅供决策参考，最终选择权在您。')
    
    doc.save(TEMPLATE_PATH)